from flask import render_template, redirect, url_for, flash, request, jsonify, send_file
from flask_login import login_required, current_user
from datetime import date, datetime
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app import db
from app.models import (Prefactura, PrefacturaDetalle, ReciboCaja, AtencionDetalle,
                        AtencionEncabezado, Cliente, FormaPago, NumeracionInicial)
from app.cartera import cartera_bp


def _generar_nro_prefactura():
    anio = date.today().year
    num = NumeracionInicial.query.filter_by(anio=anio).first()
    if not num:
        num = NumeracionInicial(anio=anio)
        db.session.add(num)
        db.session.flush()
    nro = f'{num.prefijo_prefactura}{anio}{num.consecutivo_prefactura:05d}'
    num.consecutivo_prefactura += 1
    return nro


def _generar_nro_recibo():
    anio = date.today().year
    num = NumeracionInicial.query.filter_by(anio=anio).first()
    if not num:
        num = NumeracionInicial(anio=anio)
        db.session.add(num)
        db.session.flush()
    nro = f'{num.prefijo_recibo}{anio}{num.consecutivo_recibo:05d}'
    num.consecutivo_recibo += 1
    return nro


def _formatear_documento(documento):
    texto = str(documento or '').strip()
    if texto.isdigit():
        return f'{int(texto):,}'
    return texto


def _fecha_edad_texto(paciente):
    if not paciente or not paciente.fecha_nacimiento:
        return ''
    return f'({paciente.fecha_nacimiento:%Y-%m-%d})'


def _build_prefactura_rows(prefactura):
    filas_detalladas = []
    filas_agrupadas = []
    total_prefactura = 0
    cantidad_solicitudes = 0

    detalles_prefactura = prefactura.atenciones.order_by(PrefacturaDetalle.nro_consecutivo.asc()).all()
    for item in detalles_prefactura:
        atencion = item.atencion
        if not atencion or not atencion.paciente:
            continue

        paciente = atencion.paciente
        detalles = atencion.detalles.order_by(AtencionDetalle.secuencia.asc()).all()

        total_atencion = 0
        for detalle in detalles:
            valor = int(round(float(detalle.valor or 0)))
            total_atencion += valor
            filas_detalladas.append([
                atencion.fecha_atencion,
                atencion.nro_consecutivo,
                paciente.tipo_id.codigo if paciente.tipo_id else '',
                _formatear_documento(paciente.identificacion),
                _fecha_edad_texto(paciente),
                paciente.nombre_completo,
                detalle.cups_codigo,
                detalle.cups.descripcion if detalle.cups else '',
                valor
            ])

        filas_detalladas.append([
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            'TOTAL',
            total_atencion
        ])
        filas_detalladas.append([None] * 9)

        filas_agrupadas.append([
            atencion.fecha_atencion,
            atencion.nro_consecutivo,
            paciente.tipo_id.codigo if paciente.tipo_id else '',
            _formatear_documento(paciente.identificacion),
            _fecha_edad_texto(paciente),
            paciente.nombre_completo,
            None,
            'TOTAL A PAGAR',
            total_atencion
        ])
        total_prefactura += total_atencion
        cantidad_solicitudes += 1

    if filas_detalladas and filas_detalladas[-1] == [None] * 9:
        filas_detalladas.pop()

    return filas_detalladas, filas_agrupadas, total_prefactura, cantidad_solicitudes


def _agregar_resumen_prefactura(ws, cantidad_solicitudes, total_prefactura):
    ws.append([None] * 9)
    ws.append(['CANTIDAD DE SOLICITUDES', None, None, None, None, None, None, None, cantidad_solicitudes])
    ws.append(['TOTAL A PAGAR A ZERBIT', None, None, None, None, None, None, None, total_prefactura])


def _aplicar_estilos_prefactura(ws, compacta=False):
    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(fill_type='solid', fgColor='FF00FFFF')
    total_fill = PatternFill(fill_type='solid', fgColor='FFF2CC')

    anchos = {
        'A': 12, 'B': 23.14, 'C': 11.28, 'D': 13, 'E': 12.71,
        'F': 39.57, 'G': 11.86, 'H': 25.86, 'I': 16
    }
    if compacta:
        anchos.update({'A': 10.86, 'B': 14, 'C': 11.28, 'D': 12.71, 'E': 11.86})
    for col, width in anchos.items():
        ws.column_dimensions[col].width = width

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=9):
        for cell in row:
            if any(c.value is not None for c in row):
                cell.border = border
            if cell.column == 1 and isinstance(cell.value, (date, datetime)):
                cell.number_format = 'yyyy-mm-dd;@'
            elif cell.column == 9 and cell.value is not None:
                cell.number_format = '#,##0'
        if ws.cell(row[0].row, 8).value in {'TOTAL A PAGAR', 'TOTAL'}:
            ws.cell(row[0].row, 8).fill = total_fill
            ws.cell(row[0].row, 8).font = Font(bold=True)
            ws.cell(row[0].row, 9).fill = total_fill
            ws.cell(row[0].row, 9).font = Font(bold=True)

    for row_num in range(2, ws.max_row + 1):
        etiqueta = ws[f'A{row_num}'].value
        if etiqueta in {'CANTIDAD DE SOLICITUDES', 'TOTAL A PAGAR A ZERBIT'}:
            ws.merge_cells(f'A{row_num}:H{row_num}')
            ws[f'A{row_num}'].font = Font(bold=True)
            ws[f'A{row_num}'].alignment = Alignment(horizontal='center')
            ws[f'A{row_num}'].fill = total_fill
            ws[f'I{row_num}'].font = Font(bold=True)
            ws[f'I{row_num}'].fill = total_fill
            ws[f'I{row_num}'].border = border
            if etiqueta == 'TOTAL A PAGAR A ZERBIT':
                ws[f'I{row_num}'].number_format = '"$" #,##0'


def _generar_excel_prefactura(prefactura):
    wb = Workbook()
    ws_detallado = wb.active
    ws_detallado.title = 'detallado'
    ws_agrupado = wb.create_sheet('agrupado')

    encabezados = [
        'FECHA DE LA TOMA', '# SOLICITUD', 'CEDULA/PA/CE ETC', 'N.DOCUMENTO',
        'EDAD', 'NOMBRES Y APELLIDOS', 'CUP', 'EXAMENES', 'VALOR POR EXAMEN'
    ]
    ws_detallado.append(encabezados)
    ws_agrupado.append(encabezados)

    filas_detalladas, filas_agrupadas, total_prefactura, cantidad_solicitudes = _build_prefactura_rows(prefactura)
    for fila in filas_detalladas:
        ws_detallado.append(fila)
    _agregar_resumen_prefactura(ws_detallado, cantidad_solicitudes, total_prefactura)

    for fila in filas_agrupadas:
        ws_agrupado.append(fila)
    _agregar_resumen_prefactura(ws_agrupado, cantidad_solicitudes, total_prefactura)

    _aplicar_estilos_prefactura(ws_detallado, compacta=False)
    _aplicar_estilos_prefactura(ws_agrupado, compacta=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@cartera_bp.route('/prefacturas')
@login_required
def prefacturas():
    q = request.args.get('q', '').strip()
    estado = request.args.get('estado', '')
    query = Prefactura.query
    if q:
        query = query.join(Cliente).filter(
            db.or_(
                Prefactura.nro_prefactura.ilike(f'%{q}%'),
                Cliente.nombre.ilike(f'%{q}%')
            )
        )
    if estado:
        query = query.filter_by(estado=estado)
    lista = query.order_by(Prefactura.fecha_generacion.desc()).limit(200).all()
    return render_template('cartera/prefacturas.html', prefacturas=lista, q=q, estado=estado)


@cartera_bp.route('/prefacturas/nueva', methods=['GET', 'POST'])
@login_required
def prefactura_nueva():
    clientes = Cliente.query.filter_by(estado=True).order_by(Cliente.nombre).all()

    if request.method == 'POST':
        cliente_id = request.form.get('cliente_id')
        fecha_desde_str = request.form.get('fecha_desde', '')
        fecha_hasta_str = request.form.get('fecha_hasta', '')
        atenciones_ids = request.form.getlist('atenciones_sel')

        if not cliente_id or not atenciones_ids:
            flash('Debe seleccionar el cliente y al menos una atención.', 'danger')
            return redirect(url_for('cartera.prefactura_nueva'))

        try:
            fecha_desde = datetime.strptime(fecha_desde_str, '%Y-%m-%d').date() if fecha_desde_str else None
            fecha_hasta = datetime.strptime(fecha_hasta_str, '%Y-%m-%d').date() if fecha_hasta_str else None
        except ValueError:
            flash('Fechas inválidas.', 'danger')
            return redirect(url_for('cartera.prefactura_nueva'))

        nro = _generar_nro_prefactura()
        valor_total = 0
        pf = Prefactura(
            nro_prefactura=nro,
            cliente_id=cliente_id,
            fecha_desde=fecha_desde,
            fecha_hasta=fecha_hasta,
            estado='Pendiente',
            usuario_genera_id=current_user.id
        )
        db.session.add(pf)
        db.session.flush()

        for consec in atenciones_ids:
            a = AtencionEncabezado.query.get(consec)
            if a and a.estado == 'Ingresado':
                pfd = PrefacturaDetalle(prefactura_id=pf.id, nro_consecutivo=consec)
                db.session.add(pfd)
                a.estado = 'Prefacturado'
                valor_total += float(a.valor_total)

        pf.valor_total = valor_total
        db.session.commit()
        flash(f'Prefactura {nro} generada por ${valor_total:,.0f}.', 'success')
        return redirect(url_for('cartera.prefactura_ver', id=pf.id, descargar_excel=1))

    return render_template('cartera/prefactura_nueva.html', clientes=clientes)


@cartera_bp.route('/prefacturas/<int:id>')
@login_required
def prefactura_ver(id):
    pf = Prefactura.query.get_or_404(id)
    descargar_excel = request.args.get('descargar_excel') == '1'
    total_pagado = sum(float(r.valor_pagado or 0) for r in pf.recibos)
    return render_template('cartera/prefactura_ver.html', prefactura=pf,
                           descargar_excel=descargar_excel,
                           total_pagado=total_pagado)


@cartera_bp.route('/prefacturas/<int:id>/excel')
@login_required
def prefactura_excel(id):
    pf = Prefactura.query.get_or_404(id)
    buf = _generar_excel_prefactura(pf)
    filename = f'Informe-prefactura-{pf.nro_prefactura}.xlsx'
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


@cartera_bp.route('/recibos/nuevo/<int:prefactura_id>', methods=['GET', 'POST'])
@login_required
def recibo_nuevo(prefactura_id):
    pf = Prefactura.query.get_or_404(prefactura_id)
    formas_pago = FormaPago.query.filter_by(estado=True).order_by(FormaPago.nombre).all()

    if request.method == 'POST':
        valor_pagado = float(request.form.get('valor_pagado', 0) or 0)
        forma_pago_id = request.form.get('forma_pago_id')
        observaciones = request.form.get('observaciones', '').strip()

        nro = _generar_nro_recibo()
        recibo = ReciboCaja(
            nro_recibo=nro,
            prefactura_id=pf.id,
            cliente_id=pf.cliente_id,
            valor_pagado=valor_pagado,
            forma_pago_id=forma_pago_id,
            observaciones=observaciones,
            usuario_registra_id=current_user.id
        )
        db.session.add(recibo)

        # Actualizar estado prefactura
        total_pagado = sum(r.valor_pagado for r in pf.recibos) + valor_pagado
        if total_pagado >= float(pf.valor_total):
            pf.estado = 'Pagada'
        else:
            pf.estado = 'Parcial'

        db.session.commit()
        flash(f'Recibo {nro} registrado por ${valor_pagado:,.0f}.', 'success')
        return redirect(url_for('cartera.prefactura_ver', id=pf.id))

    return render_template('cartera/recibo_nuevo.html', prefactura=pf, formas_pago=formas_pago)


@cartera_bp.route('/movimiento')
@login_required
def movimiento():
    fecha_desde = request.args.get('fecha_desde', '')
    fecha_hasta = request.args.get('fecha_hasta', '')
    cliente_id = request.args.get('cliente_id', '')
    clientes = Cliente.query.filter_by(estado=True).order_by(Cliente.nombre).all()
    resultados = []

    if fecha_desde and fecha_hasta:
        try:
            fd = datetime.strptime(fecha_desde, '%Y-%m-%d').date()
            fh = datetime.strptime(fecha_hasta, '%Y-%m-%d').date()
            query = Prefactura.query.filter(
                db.func.date(Prefactura.fecha_generacion) >= fd,
                db.func.date(Prefactura.fecha_generacion) <= fh
            )
            if cliente_id:
                query = query.filter_by(cliente_id=cliente_id)
            resultados = query.order_by(Prefactura.fecha_generacion).all()
        except ValueError:
            flash('Fechas inválidas.', 'danger')

    return render_template('cartera/movimiento.html',
                           resultados=resultados, clientes=clientes,
                           fecha_desde=fecha_desde, fecha_hasta=fecha_hasta,
                           cliente_id=cliente_id)


@cartera_bp.route('/recaudos')
@login_required
def recaudos():
    fecha_desde = request.args.get('fecha_desde', '')
    fecha_hasta = request.args.get('fecha_hasta', '')
    clientes = Cliente.query.filter_by(estado=True).order_by(Cliente.nombre).all()
    resultados = []

    if fecha_desde and fecha_hasta:
        try:
            fd = datetime.strptime(fecha_desde, '%Y-%m-%d').date()
            fh = datetime.strptime(fecha_hasta, '%Y-%m-%d').date()
            resultados = ReciboCaja.query.filter(
                db.func.date(ReciboCaja.fecha_pago) >= fd,
                db.func.date(ReciboCaja.fecha_pago) <= fh
            ).order_by(ReciboCaja.fecha_pago).all()
        except ValueError:
            flash('Fechas inválidas.', 'danger')

    return render_template('cartera/recaudos.html', resultados=resultados,
                           clientes=clientes, fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)


@cartera_bp.route('/estado-cuenta')
@login_required
def estado_cuenta():
    cliente_id = request.args.get('cliente_id', '')
    clientes = Cliente.query.filter_by(estado=True).order_by(Cliente.nombre).all()
    cliente = None
    prefacturas_lista = []

    if cliente_id:
        cliente = Cliente.query.get(cliente_id)
        prefacturas_lista = Prefactura.query.filter_by(
            cliente_id=cliente_id
        ).order_by(Prefactura.fecha_generacion).all()

    return render_template('cartera/estado_cuenta.html',
                           clientes=clientes, cliente=cliente,
                           prefacturas=prefacturas_lista, cliente_id=cliente_id)


@cartera_bp.route('/api/atenciones-pendientes/<int:cliente_id>')
@login_required
def api_atenciones_pendientes(cliente_id):
    atenciones = AtencionEncabezado.query.filter_by(
        cliente_id=cliente_id, estado='Ingresado'
    )

    fecha_desde = request.args.get('fecha_desde', '').strip()
    fecha_hasta = request.args.get('fecha_hasta', '').strip()
    try:
        if fecha_desde:
            atenciones = atenciones.filter(AtencionEncabezado.fecha_atencion >= datetime.strptime(fecha_desde, '%Y-%m-%d').date())
        if fecha_hasta:
            atenciones = atenciones.filter(AtencionEncabezado.fecha_atencion <= datetime.strptime(fecha_hasta, '%Y-%m-%d').date())
    except ValueError:
        return jsonify({'error': 'Fechas inválidas.'}), 400

    atenciones = atenciones.order_by(AtencionEncabezado.fecha_atencion).all()
    return jsonify([{
        'consecutivo': a.nro_consecutivo,
        'paciente': a.paciente.nombre_completo if a.paciente else '',
        'fecha': str(a.fecha_atencion),
        'valor_total': float(a.valor_total)
    } for a in atenciones])
