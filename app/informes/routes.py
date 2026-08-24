from flask import render_template, request, send_file
from flask_login import login_required
from app import db
from app.models import (AtencionEncabezado, AtencionDetalle, Cliente, Paciente,
                        PersonaAtiende)
from app.informes import informes_bp
from datetime import datetime
import io


@informes_bp.route('/')
@login_required
def index():
    return render_template('informes/index.html')


@informes_bp.route('/analisis-empresa')
@login_required
def analisis_empresa():
    """Análisis de servicios prestados agrupados por empresa"""
    fecha_desde = request.args.get('fecha_desde', '')
    fecha_hasta = request.args.get('fecha_hasta', '')
    clientes = Cliente.query.filter_by(estado=True).order_by(Cliente.nombre).all()
    datos = []

    if fecha_desde and fecha_hasta:
        try:
            fd = datetime.strptime(fecha_desde, '%Y-%m-%d').date()
            fh = datetime.strptime(fecha_hasta, '%Y-%m-%d').date()
            # Agrupar atenciones por cliente
            from sqlalchemy import func
            resultado = db.session.query(
                Cliente.nombre,
                func.count(AtencionEncabezado.nro_consecutivo).label('total_atenciones'),
                func.sum(AtencionDetalle.valor).label('total_valor')
            ).join(AtencionEncabezado, AtencionEncabezado.cliente_id == Cliente.id
            ).join(AtencionDetalle, AtencionDetalle.nro_consecutivo == AtencionEncabezado.nro_consecutivo
            ).filter(
                AtencionEncabezado.fecha_atencion >= fd,
                AtencionEncabezado.fecha_atencion <= fh
            ).group_by(Cliente.nombre).all()
            datos = resultado
        except ValueError:
            pass

    return render_template('informes/analisis_empresa.html',
                           datos=datos, clientes=clientes,
                           fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)


@informes_bp.route('/analisis-empresa/excel')
@login_required
def analisis_empresa_excel():
    import pandas as pd
    fecha_desde = request.args.get('fecha_desde', '')
    fecha_hasta = request.args.get('fecha_hasta', '')
    datos = []

    if fecha_desde and fecha_hasta:
        try:
            fd = datetime.strptime(fecha_desde, '%Y-%m-%d').date()
            fh = datetime.strptime(fecha_hasta, '%Y-%m-%d').date()
            from sqlalchemy import func
            resultado = db.session.query(
                Cliente.nombre,
                func.count(AtencionEncabezado.nro_consecutivo).label('total_atenciones'),
                func.sum(AtencionDetalle.valor).label('total_valor')
            ).join(AtencionEncabezado, AtencionEncabezado.cliente_id == Cliente.id
            ).join(AtencionDetalle, AtencionDetalle.nro_consecutivo == AtencionEncabezado.nro_consecutivo
            ).filter(
                AtencionEncabezado.fecha_atencion >= fd,
                AtencionEncabezado.fecha_atencion <= fh
            ).group_by(Cliente.nombre).all()
            datos = [{'Empresa': r[0], 'Atenciones': r[1], 'Valor Total': float(r[2] or 0)}
                     for r in resultado]
        except ValueError:
            pass

    df = pd.DataFrame(datos) if datos else pd.DataFrame(columns=['Empresa', 'Atenciones', 'Valor Total'])
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Análisis Empresa')
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'analisis_empresa_{fecha_desde}_{fecha_hasta}.xlsx')


@informes_bp.route('/remision-excel')
@login_required
def remision_excel():
    """Genera Excel de remisión agrupado por persona que atiende"""
    fecha_desde = request.args.get('fecha_desde', '')
    fecha_hasta = request.args.get('fecha_hasta', '')
    persona_id = request.args.get('persona_id', '')
    personas = PersonaAtiende.query.filter_by(estado=True).order_by(
        PersonaAtiende.apellidos, PersonaAtiende.nombres).all()

    if not fecha_desde or not fecha_hasta:
        return render_template('informes/remision.html',
                               personas=personas, datos=None,
                               fecha_desde=fecha_desde, fecha_hasta=fecha_hasta,
                               persona_id=persona_id)

    try:
        fd = datetime.strptime(fecha_desde, '%Y-%m-%d').date()
        fh = datetime.strptime(fecha_hasta, '%Y-%m-%d').date()
    except ValueError:
        return render_template('informes/remision.html',
                               personas=personas, datos=None,
                               fecha_desde=fecha_desde, fecha_hasta=fecha_hasta,
                               persona_id=persona_id)

    query = AtencionEncabezado.query.filter(
        AtencionEncabezado.fecha_atencion >= fd,
        AtencionEncabezado.fecha_atencion <= fh,
        AtencionEncabezado.estado != 'Anulado'
    )
    if persona_id:
        query = query.filter_by(persona_atiende_id=persona_id)

    atenciones = query.order_by(
        AtencionEncabezado.persona_atiende_id,
        AtencionEncabezado.fecha_atencion
    ).all()

    # Si el usuario solo quiere ver preview, mostramos tabla
    if request.args.get('formato') != 'excel':
        return render_template('informes/remision.html',
                               personas=personas, datos=atenciones,
                               fecha_desde=fecha_desde, fecha_hasta=fecha_hasta,
                               persona_id=persona_id)

    # Generar Excel estilo remisión
    import pandas as pd
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, Alignment

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        # Agrupar por persona que atiende
        grupos = {}
        for a in atenciones:
            persona_nombre = a.persona_atiende.nombre_completo if a.persona_atiende else 'Sin asignar'
            if persona_nombre not in grupos:
                grupos[persona_nombre] = []
            # Un registro por atención con todos los exámenes concatenados
            examenes = ', '.join(d.cups.descripcion for d in a.detalles if d.cups)
            muestras = ''  # Se puede llenar después si se agrega tipo de muestra
            grupos[persona_nombre].append({
                'NOMBRES Y APELLIDOS DEL PACIENTE': a.paciente.nombre_completo if a.paciente else '',
                'FECHA DE NACIMIENTO': a.paciente.fecha_nacimiento if a.paciente else '',
                'No DE DOCUMENTO': a.paciente.identificacion if a.paciente else '',
                'CIUDAD': (
                    'BOGOTA' if a.paciente and a.paciente.municipio and a.paciente.municipio.es_localidad
                    else (a.paciente.departamento.nombre if a.paciente and a.paciente.departamento
                          else 'BOGOTA')
                ) if a.paciente else 'BOGOTA',
                'GENERO': a.paciente.sexo.codigo if a.paciente and a.paciente.sexo else '',
                'DIAGNOSTICO': '',
                'EXAMEN SOLICITADO': examenes,
            })

        # Crear una hoja por persona o todo en una hoja
        fila_actual = 0
        sheet_name = 'Remisión'
        all_rows = []
        for persona, registros in grupos.items():
            # Header de persona
            all_rows.append({'NOMBRES Y APELLIDOS DEL PACIENTE': f'--- PERSONA QUE ENVÍA: {persona} ---'})
            for r in registros:
                all_rows.append(r)
            all_rows.append({})  # fila vacía separador

        df = pd.DataFrame(all_rows)
        df.to_excel(writer, index=False, sheet_name=sheet_name, startrow=5)

        # Agregar encabezado
        ws = writer.sheets[sheet_name]
        ws['A1'] = 'LABORATORIO GENERAL Y ESPECIALIZADO'
        ws['A1'].font = Font(bold=True, size=12)
        ws['A2'] = f'SOLICITUD DE EXAMENES - REMISION'
        ws['A3'] = f'CIUDAD Y FECHA: BOGOTA {fd.strftime("%d/%m/%Y")} al {fh.strftime("%d/%m/%Y")}'
        ws['A4'] = f'LABORATORIO REMITENTE: ZERBIT SAS'

        # Ajustar ancho de columnas
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(20, len(str(col)) + 5)

    buf.seek(0)
    filename = f'Remision_{fecha_desde}_{fecha_hasta}.xlsx'
    return send_file(buf,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)
